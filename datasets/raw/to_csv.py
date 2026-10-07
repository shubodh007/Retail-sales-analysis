"""One-off UCI xlsx -> canonical CSV. Uses system openpyxl (not a project dep).

Reads both year sheets, normalizes headers to the canonical contract names,
writes datasets/raw/online_retail_II.csv (UTF-8). Deterministic: sheet order
fixed, no sampling. Run once; the CSV is the pipeline input.
"""
import csv
import sys
from pathlib import Path

RAW = Path(__file__).resolve().parent
XLSX = RAW / "online_retail_II.xlsx"
OUT = RAW / "online_retail_II.csv"
SHEETS = ["Year 2009-2010", "Year 2010-2011"]
HEADER = ["InvoiceNo", "StockCode", "Description", "Quantity", "InvoiceDate",
          "UnitPrice", "CustomerID", "Country"]

try:
    from openpyxl import load_workbook
except ImportError:
    sys.exit("openpyxl missing: use the system python that has it")


def main() -> None:
    wb = load_workbook(filename=str(XLSX), read_only=True, data_only=True)
    print("sheets:", wb.sheetnames)
    total = 0
    with open(OUT, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(HEADER)
        for name in SHEETS:
            ws = wb[name]
            rows = ws.iter_rows(values_only=True)
            header = next(rows)  # skip source header
            print(name, "header:", list(header))
            n = 0
            for r in rows:
                if r[0] is None and r[3] is None:
                    continue
                w.writerow(list(r[:8]))
                n += 1
            print(name, "rows:", n)
            total += n
    print("wrote", OUT, total, "rows")


if __name__ == "__main__":
    main()
